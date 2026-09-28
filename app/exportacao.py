"""Relatório de chamados em Excel (openpyxl) e PDF (reportlab), gerados em memória."""

import io
from datetime import date, datetime
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import formatos

# (cabeçalho, valor, largura Excel em caracteres, largura PDF em mm)
COLUNAS = [
    ("Nº", lambda c: c.id, 8, 13),
    ("Título", lambda c: c.titulo, 40, 50),
    ("Fila", lambda c: c.fila.nome, 16, 24),
    ("Tipo", lambda c: c.categoria.nome if c.categoria else "", 18, 24),
    ("Status", lambda c: c.status, 18, 24),
    ("Prioridade", lambda c: c.prioridade, 11, 17),
    ("Setor de origem", lambda c: c.setor, 16, 22),
    ("Solicitante", lambda c: c.solicitante.nome, 22, 25),
    ("Responsável", lambda c: c.responsavel.nome if c.responsavel else "", 22, 25),
    ("Abertura", lambda c: formatos.data(c.aberto_em), 12, 19),
    ("Fechamento", lambda c: formatos.data(c.fechado_em), 12, 19),
    ("Avaliação", lambda c: c.avaliacao or "", 11, 15),
]
AZUL = "1F3A5F"


def gerar_excel(chamados: list) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Chamados"
    fundo, fonte = PatternFill("solid", fgColor=AZUL), Font(bold=True, color="FFFFFF")
    for col, (titulo, _c, largura, _p) in enumerate(COLUNAS, start=1):
        cel = ws.cell(row=1, column=col, value=titulo)
        cel.fill, cel.font = fundo, fonte
        ws.column_dimensions[cel.column_letter].width = largura
    for c in chamados:
        ws.append([valor(c) for _t, valor, _l, _p in COLUNAS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def gerar_pdf(chamados: list, inicio: date, fim: date) -> io.BytesIO:
    buf = io.BytesIO()
    pagina = landscape(A4)
    doc = SimpleDocTemplate(buf, pagesize=pagina, leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title="Relatório de Chamados")
    estilos = getSampleStyleSheet()
    celula = estilos["BodyText"].clone("celula", fontSize=8, leading=10)

    dados = [[t for t, *_ in COLUNAS]]
    for c in chamados:
        dados.append([Paragraph(escape(str(valor(c))), celula) for _t, valor, _l, _p in COLUNAS])

    tabela = Table(dados, colWidths=[p * mm for *_x, p in COLUNAS], repeatRows=1)
    tabela.hAlign = "LEFT"
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + AZUL)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F8")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C8CDD5")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))

    def rodape(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(10 * mm, 7 * mm, f"Gerado em {datetime.now():%d/%m/%Y %H:%M}")
        canvas.drawRightString(pagina[0] - 10 * mm, 7 * mm, f"Página {doc_.page}")
        canvas.restoreState()

    doc.build([
        Paragraph("Relatório de Chamados", estilos["Title"]),
        Paragraph(f"Período: {inicio:%d/%m/%Y} a {fim:%d/%m/%Y} — {len(chamados)} chamado(s)",
                  estilos["Normal"]),
        Spacer(1, 6 * mm),
        tabela,
    ], onFirstPage=rodape, onLaterPages=rodape)
    buf.seek(0)
    return buf
