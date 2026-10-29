"""Testes das rotas, rodando num banco SQLite em memória.

    python -m pytest tests
"""

import io
import os
import re

import pytest
from sqlalchemy import func, select

from app import servicos
from app.modelos import Anexo, Chamado, Comentario, Usuario
from tests.conftest import csrf, entrar


def chamado(s, titulo) -> Chamado:
    s.expire_all()
    return s.scalar(select(Chamado).where(Chamado.titulo == titulo))


def usuario(s, login) -> Usuario:
    s.expire_all()
    return s.scalar(select(Usuario).where(Usuario.login == login))


def abrir(client, s, titulo, setor="T.I", para="carlos", **extra):
    """Abre um chamado pelo formulário e devolve a resposta."""
    dados = {"csrf": csrf(client, "/chamados/novo"), "setor_destino": setor,
             "responsavel": usuario(s, para).id if para else "",
             "titulo": titulo, "descricao": "Detalhes", "prioridade": "Baixa"}
    dados.update(extra)
    return client.post("/chamados/novo", content_type="multipart/form-data", data=dados)


def sair(client):
    client.post("/logout", data={"csrf": csrf(client, "/conta/")})


def salvar_usuarios(client, s, linhas: dict):
    """Salva pela tela de edição de cada conta: {login: (setor, atende, papel)}."""
    r = None
    for login, (setor, atende, papel) in linhas.items():
        u = usuario(s, login)
        dados = {"csrf": csrf(client, f"/usuarios/{u.id}/editar"), "nome": u.nome, "login": u.login,
                 "email": u.email or "", "setor": setor, "papel": papel}
        if atende:
            dados["atende"] = "1"
        r = client.post(f"/usuarios/{u.id}/editar", data=dados, follow_redirects=True)
    return r


def salvar_usuario(client, s, login, setor, atende, papel="usuario"):
    return salvar_usuarios(client, s, {login: (setor, atende, papel)})


# ─── cadastro e login ───────────────────────────────────────────────────────

def cadastrar(client, **campos):
    dados = {"nome": "Maria Teste", "login": "maria", "email": "maria@empresa.com.br", "setor": "RH",
             "senha": "senha123", "confirmar": "senha123"}
    dados.update(campos)
    return client.post("/cadastro", data=dict(dados, csrf=csrf(client, "/cadastro")))


def test_primeira_conta_vira_admin_e_as_outras_usuario(client, s):
    r = cadastrar(client)
    assert r.status_code == 302
    assert usuario(s, "maria").is_admin
    sair(client)
    cadastrar(client, login="joao", nome="João")
    joao = usuario(s, "joao")
    assert joao.papel == "usuario" and not joao.atende   # ninguém atende sem o admin marcar


def test_cadastro_valida(client, contas):
    assert "já está em uso" in cadastrar(client, login="ana").get_data(as_text=True)
    assert "não conferem" in cadastrar(client, confirmar="outra123").get_data(as_text=True)
    assert "8 caracteres" in cadastrar(client, senha="abc1", confirmar="abc1").get_data(as_text=True)
    assert "letras e números" in cadastrar(client, senha="somenteletras", confirmar="somenteletras").get_data(as_text=True)
    assert "Usuário: de 3 a 50" in cadastrar(client, login="a b").get_data(as_text=True)


def test_login_normaliza_maiusculas_e_senha_fica_em_hash(client, s, contas):
    assert usuario(s, "ana").senha_hash != "ana12345"
    r = entrar(client, "ANA", "ana12345")
    assert r.status_code == 302


def test_exige_login(client):
    r = client.get("/chamados")
    assert r.status_code == 302 and "/login" in r.location


def test_login_invalido_e_bloqueio(client, contas):
    assert "inválidos" in entrar(client, "ana", "errada").get_data(as_text=True)
    for _ in range(4):
        entrar(client, "ana", "errada")
    assert "Muitas tentativas" in entrar(client, "ana", "ana12345").get_data(as_text=True)


def test_sem_csrf_recusado(client, contas):
    assert client.post("/login", data={"login": "ana", "senha": "ana12345"}).status_code == 400


def test_redirect_externo_bloqueado(client, contas):
    r = client.post("/login", data={"login": "admin", "senha": "admin123", "csrf": csrf(client),
                                    "proximo": "//site-malicioso.com"})
    assert r.location.endswith("/chamados")


def test_conta_desativada_perde_acesso_na_hora(client, s, contas):
    entrar(client, "ana", "ana12345")
    assert client.get("/chamados").status_code == 200
    u = usuario(s, "ana")
    u.ativo = False
    s.commit()
    assert client.get("/chamados").status_code == 302


def test_usuario_nao_muda_o_proprio_setor(client, s, contas):
    entrar(client, "ana", "ana12345")
    client.post("/conta/", data={"csrf": csrf(client, "/conta/"), "nome": "Ana S.",
                                 "email": "ana@empresa.com.br", "setor": "RH"})
    ana = usuario(s, "ana")
    assert ana.nome == "Ana S." and ana.setor == "Fiscal"


# ─── administração de contas ────────────────────────────────────────────────

def test_admin_define_setor_atende_e_perfil(client, s, contas):
    entrar(client, "admin", "admin123")
    salvar_usuario(client, s, "bruno", "Manutenção", True)
    bruno = usuario(s, "bruno")
    assert bruno.setor == "Manutenção" and bruno.atende and not bruno.is_admin
    salvar_usuario(client, s, "bruno", "Manutenção", True, papel="admin")
    assert usuario(s, "bruno").is_admin


def test_nao_remove_ultimo_admin(s, contas):
    with pytest.raises(servicos.ErroValidacao):
        servicos.atualizar_usuario(s, usuario(s, "admin"), "T.I", True, "usuario")
    with pytest.raises(servicos.ErroValidacao):
        servicos.definir_ativo(s, usuario(s, "admin"), False)


def test_tirar_atendimento_ou_desativar_libera_chamados(s, contas):
    servicos.atualizar_usuario(s, usuario(s, "carlos"), "T.I", False, "usuario")
    assert chamado(s, "Sem acesso à pasta da rede").responsavel_id is None
    assert chamado(s, "Trocar teclado").responsavel.login == "carlos"   # fechado: histórico fica
    servicos.definir_ativo(s, usuario(s, "marta"), False)
    assert chamado(s, "Tomada queimada na recepção").responsavel_id is None


# tela: (usuário comum, supervisor, administrador)
TELAS = {
    "/chamados": (200, 200, 200),
    "/painel": (200, 200, 200),
    "/relatorio": (403, 200, 200),
    "/documentos/": (403, 200, 200),
    "/inventario/": (403, 403, 200),
    "/usuarios": (403, 403, 200),
}


def test_telas_por_nivel_de_acesso(client, s, contas):
    """As telas dependem só do nível; "Atende chamados" não muda nada aqui."""
    marta = usuario(s, "marta")                       # atende a Manutenção, nível usuário
    for login, senha, nivel in (("ana", "ana12345", 0), ("marta", "marta123", 0)):
        entrar(client, login, senha)
        for url, esperado in TELAS.items():
            assert client.get(url).status_code == esperado[nivel], (login, url)
        menu = client.get("/chamados").get_data(as_text=True)
        assert ">Painel<" in menu and ">Relatório<" not in menu and ">Documentos<" not in menu
        sair(client)

    servicos.atualizar_usuario(s, marta, "Manutenção", False, "supervisor")    # supervisor que não atende
    entrar(client, "marta", "marta123")
    for url, esperado in TELAS.items():
        assert client.get(url).status_code == esperado[1], ("supervisor", url)
    menu = client.get("/chamados").get_data(as_text=True)
    assert ">Relatório<" in menu and ">Documentos<" in menu and ">Inventário<" not in menu
    sair(client)

    entrar(client, "admin", "admin123")
    for url, esperado in TELAS.items():
        assert client.get(url).status_code == esperado[2], ("admin", url)


def test_supervisor_ve_os_chamados_do_setor_sem_atender(client, s, contas):
    servicos.atualizar_usuario(s, usuario(s, "carlos"), "T.I", False, "supervisor")
    entrar(client, "carlos", "carlos123")
    html = client.get("/chamados").get_data(as_text=True)
    assert "Impressora do fiscal não imprime" in html and "Dúvida sobre saldo de férias" not in html
    c = chamado(s, "Impressora do fiscal não imprime")
    assert client.get(f"/chamados/{c.id}").status_code == 200
    # ver não é atender: sem "Atende chamados" não muda status nem some na lista de funcionários
    token = csrf(client, f"/chamados/{c.id}")
    assert client.post(f"/chamados/{c.id}/status", data={"csrf": token, "status": "Fechado"}).status_code == 403


# ─── abrir chamado: setor e funcionário ─────────────────────────────────────

def test_novo_chamado_lista_so_quem_atende(client, contas):
    entrar(client, "ana", "ana12345")
    html = client.get("/chamados/novo").get_data(as_text=True)
    assert "Carlos Técnico" in html and "Marta Manutenção" in html and "Rita Recursos Humanos" in html
    assert "Bruno Lima" not in html                   # não atende chamados
    assert '<option >Fiscal</option>' not in html      # setor sem ninguém atendendo não aparece


def test_funcionario_obrigatorio_e_do_setor_certo(client, s, contas):
    entrar(client, "ana", "ana12345")
    assert "Escolha um funcionário do setor T.I" in abrir(client, s, "Sem pessoa", para=None).get_data(as_text=True)
    r = abrir(client, s, "Pessoa errada", setor="T.I", para="marta")   # marta é da Manutenção
    assert "Escolha um funcionário do setor T.I" in r.get_data(as_text=True)
    r = abrir(client, s, "Não atende", setor="Vendas", para="bruno")   # bruno não atende
    assert "Escolha um funcionário do setor Vendas" in r.get_data(as_text=True)
    assert chamado(s, "Pessoa errada") is None


def test_abrir_chamado_para_outro_setor(client, s, contas):
    entrar(client, "ana", "ana12345")
    assert abrir(client, s, "Cadeira quebrada", setor="Manutenção", para="marta").status_code == 302
    c = chamado(s, "Cadeira quebrada")
    assert c.setor_destino == "Manutenção" and c.responsavel.login == "marta"
    assert c.setor == "Fiscal"                         # setor de quem pediu vem do cadastro


# ─── visibilidade ───────────────────────────────────────────────────────────

def test_usuario_ve_so_os_proprios(client, s, contas):
    entrar(client, "ana", "ana12345")
    html = client.get("/chamados").get_data(as_text=True)
    assert "Impressora do fiscal" in html and "Tomada queimada" in html   # dela, em setores diferentes
    assert "Trocar teclado" not in html
    assert client.get(f"/chamados/{chamado(s, 'Trocar teclado').id}").status_code == 403


def test_todos_que_atendem_o_setor_veem_o_chamado(client, s, contas):
    entrar(client, "admin", "admin123")               # admin também atende T.I
    salvar_usuario(client, s, "admin", "T.I", True, papel="admin")
    sair(client)
    entrar(client, "carlos", "carlos123")
    c = chamado(s, "Erro ao emitir nota fiscal")      # responsável é o admin
    assert client.get(f"/chamados/{c.id}").status_code == 200           # carlos, colega do setor, vê
    client.post(f"/chamados/{c.id}/assumir", data={"csrf": csrf(client, "/chamados")})
    assert chamado(s, "Erro ao emitir nota fiscal").responsavel.login == "carlos"


def test_quem_atende_ve_so_o_proprio_setor(client, s, contas):
    entrar(client, "marta", "marta123")               # Manutenção
    html = client.get("/chamados").get_data(as_text=True)
    assert "Ar-condicionado da sala" in html and "Tomada queimada" in html
    assert "Impressora do fiscal" not in html and "Dúvida sobre saldo" not in html
    c = chamado(s, "Impressora do fiscal não imprime")
    assert client.get(f"/chamados/{c.id}").status_code == 403
    r = client.post(f"/chamados/{c.id}/status", data={"csrf": csrf(client, "/chamados"), "status": "Fechado"})
    assert r.status_code == 403 and chamado(s, "Impressora do fiscal não imprime").status == "Aberto"


def test_admin_ve_todos(client, contas):
    entrar(client, "admin", "admin123")
    html = client.get("/chamados").get_data(as_text=True)
    assert "Impressora do fiscal" in html and "Tomada queimada" in html and "Dúvida sobre saldo" in html
    html = client.get("/chamados?setor=RH").get_data(as_text=True)
    assert "Dúvida sobre saldo" in html and "Impressora do fiscal" not in html


def test_quem_atende_e_pede_para_outro_setor_ve_como_solicitante(client, s, contas):
    entrar(client, "marta", "marta123")
    abrir(client, s, "Computador da manutenção travando", setor="T.I", para="carlos")
    c = chamado(s, "Computador da manutenção travando")
    assert client.get(f"/chamados/{c.id}").status_code == 200           # vê: é dela
    r = client.post(f"/chamados/{c.id}/assumir", data={"csrf": csrf(client, "/chamados")})
    assert r.status_code == 403                                         # mas não atende T.I


def test_visoes_e_busca(client, contas):
    entrar(client, "carlos", "carlos123")
    html = client.get("/chamados?visao=atribuidos&status=abertos").get_data(as_text=True)
    assert "Sem acesso à pasta" in html and "Erro ao emitir nota" not in html
    html = client.get("/chamados?visao=setor").get_data(as_text=True)
    assert "Erro ao emitir nota" in html and "Tomada queimada" not in html
    html = client.get("/chamados?visao=meus").get_data(as_text=True)
    assert "Nenhum chamado com esses filtros" in html                  # carlos não abriu nenhum
    html = client.get("/chamados?q=NOTA").get_data(as_text=True)
    assert "Erro ao emitir nota" in html and "Impressora" not in html
    html = client.get("/chamados?q=%25").get_data(as_text=True)        # "%" não casa com tudo
    assert "Nenhum chamado com esses filtros" in html
    for coluna in ("titulo", "para", "status", "abertura", "responsavel", "prioridade"):
        for direcao in ("asc", "desc"):
            assert client.get(f"/chamados?ordenar={coluna}&direcao={direcao}").status_code == 200


# ─── atendimento ────────────────────────────────────────────────────────────

def test_fluxo_completo(client, s, app, contas):
    entrar(client, "ana", "ana12345")
    r = abrir(client, s, "Mouse quebrado", anexos=(io.BytesIO(b"\x89PNG foto"), "foto.png"))
    assert r.status_code == 302
    c = chamado(s, "Mouse quebrado")
    assert c.status == "Aberto" and c.solicitante.login == "ana" and c.responsavel.login == "carlos"
    anexo = c.anexos[0]
    assert os.path.exists(os.path.join(app.config["ANEXOS_PASTA"], anexo.arquivo))
    assert client.get(f"/anexos/{anexo.id}").data == b"\x89PNG foto"

    token = csrf(client, "/chamados")
    client.post(f"/chamados/{c.id}/comentar", content_type="multipart/form-data",
                data={"csrf": token, "texto": "Urgente!", "anexos": (io.BytesIO(b"log"), "erro.log")})
    c = chamado(s, "Mouse quebrado")
    assert c.comentarios[-1].texto == "Urgente!" and c.comentarios[-1].anexos[0].nome == "erro.log"
    assert client.post(f"/chamados/{c.id}/status", data={"csrf": token, "status": "Fechado"}).status_code == 403

    # carlos atende e fecha
    sair(client)
    entrar(client, "carlos", "carlos123")
    token = csrf(client, f"/chamados/{c.id}")
    client.post(f"/chamados/{c.id}/status", data={"csrf": token, "status": "Em andamento"})
    client.post(f"/chamados/{c.id}/status", data={"csrf": token, "status": "Fechado"})
    c = chamado(s, "Mouse quebrado")
    assert c.fechado_em is not None
    assert [m.texto for m in c.comentarios if m.evento] == [
        'Status alterado de "Aberto" para "Em andamento".',
        'Status alterado de "Em andamento" para "Fechado".',
    ]

    # outro usuário não baixa o anexo nem avalia
    sair(client)
    entrar(client, "bruno", "bruno123")
    assert client.get(f"/anexos/{anexo.id}").status_code == 403
    client.post(f"/chamados/{c.id}/avaliar", data={"csrf": csrf(client, "/chamados"), "avaliacao": "Bom"})
    assert chamado(s, "Mouse quebrado").avaliacao is None

    # a dona avalia
    sair(client)
    entrar(client, "ana", "ana12345")
    client.post(f"/chamados/{c.id}/avaliar", data={"csrf": csrf(client, "/avaliacoes"), "avaliacao": "Bom"})
    assert chamado(s, "Mouse quebrado").avaliacao == "Bom"


def test_encaminhar_para_outro_setor(client, s, contas):
    entrar(client, "carlos", "carlos123")
    c = chamado(s, "Sem acesso à pasta da rede")                  # T.I, com carlos
    r = client.post(f"/chamados/{c.id}/encaminhar",
                    data={"csrf": csrf(client, "/chamados"), "setor_destino": "Manutenção",
                          "responsavel": usuario(s, "marta").id})
    assert r.location.endswith("/chamados")                      # carlos não atende Manutenção: volta à lista
    c = chamado(s, "Sem acesso à pasta da rede")
    assert c.setor_destino == "Manutenção" and c.responsavel.login == "marta"
    assert "Encaminhado de T.I para Manutenção (Marta Manutenção)." in [m.texto for m in c.comentarios]
    assert client.get(f"/chamados/{c.id}").status_code == 403    # e deixa de ver
    sair(client)
    entrar(client, "marta", "marta123")
    assert client.get(f"/chamados/{c.id}").status_code == 200


def test_encaminhar_exige_funcionario_do_setor(client, s, contas):
    entrar(client, "carlos", "carlos123")
    c = chamado(s, "Impressora do fiscal não imprime")
    for setor, pessoa in (("Manutenção", "carlos"), ("T.I", "bruno"), ("T.I", "marta")):
        html = client.post(f"/chamados/{c.id}/encaminhar", follow_redirects=True,
                           data={"csrf": csrf(client, "/chamados"), "setor_destino": setor,
                                 "responsavel": usuario(s, pessoa).id}).get_data(as_text=True)
        assert f"Escolha um funcionário do setor {setor}" in html
    c = chamado(s, "Impressora do fiscal não imprime")
    assert c.setor_destino == "T.I" and c.responsavel.login == "carlos"


def test_reabrir_limpa_fechamento_e_avaliacao(client, s, contas):
    entrar(client, "carlos", "carlos123")
    c = chamado(s, "E-mail não sincroniza no celular")
    assert c.avaliacao == "Bom"
    client.post(f"/chamados/{c.id}/status", data={"csrf": csrf(client, "/chamados"), "status": "Em andamento"})
    c = chamado(s, "E-mail não sincroniza no celular")
    assert c.fechado_em is None and c.avaliacao is None


def test_so_admin_apaga(client, s, app, contas):
    entrar(client, "ana", "ana12345")
    abrir(client, s, "Apagar-me", anexos=(io.BytesIO(b"abc"), "a.txt"))
    c = chamado(s, "Apagar-me")
    arquivo = os.path.join(app.config["ANEXOS_PASTA"], c.anexos[0].arquivo)
    cid = c.id
    sair(client)

    entrar(client, "carlos", "carlos123")                      # atende o setor: não apaga
    assert client.post(f"/chamados/{cid}/apagar", data={"csrf": csrf(client, "/chamados")}).status_code == 403
    sair(client)

    entrar(client, "admin", "admin123")
    client.post(f"/chamados/{cid}/apagar", data={"csrf": csrf(client, "/chamados")})
    s.expire_all()
    assert s.get(Chamado, cid) is None
    assert not s.scalars(select(Comentario).where(Comentario.chamado_id == cid)).all()
    assert not s.scalars(select(Anexo).where(Anexo.chamado_id == cid)).all()
    assert not os.path.exists(arquivo)


# ─── painel e relatório ─────────────────────────────────────────────────────

def test_painel_relatorio_por_setor(client, s, contas):
    servicos.atualizar_usuario(s, usuario(s, "marta"), "Manutenção", True, "supervisor")
    entrar(client, "marta", "marta123")
    assert client.get("/painel").status_code == 200
    r = client.get("/relatorio/excel?de=2000-01-01&ate=2100-01-01&setor=T.I")   # setor ignorado: não é admin
    assert r.status_code == 200 and r.data[:2] == b"PK"
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.data)).active
    assert {row[2] for row in ws.iter_rows(min_row=2, values_only=True)} == {"Manutenção"}

    sair(client)
    entrar(client, "admin", "admin123")
    for url in ("/painel", "/painel?dias=365", "/painel?setor=T.I", "/usuarios", "/relatorio", "/conta/"):
        assert client.get(url).status_code == 200, url
    r = client.get("/relatorio/pdf?de=2000-01-01&ate=2100-01-01")
    assert r.status_code == 200 and r.data[:4] == b"%PDF"
    assert client.get("/relatorio/pdf?de=2100-01-01&ate=2000-01-01").status_code == 302


# ─── anexos e segurança ─────────────────────────────────────────────────────

def test_anexo_proibido_ou_grande(client, s, contas):
    entrar(client, "ana", "ana12345")
    r = abrir(client, s, "X", anexos=(io.BytesIO(b"MZ"), "virus.exe"))
    assert "não permitido" in r.get_data(as_text=True)
    r = abrir(client, s, "X", anexos=(io.BytesIO(b"0" * (1024 * 1024 + 1)), "grande.txt"))
    assert "limite" in r.get_data(as_text=True)


def test_anexo_nao_exibivel_sempre_baixado(client, s, contas):
    entrar(client, "ana", "ana12345")
    abrir(client, s, "Com xml", anexos=(io.BytesIO(b"<script>alert(1)</script>"), "dados.xml", "text/html"))
    a = chamado(s, "Com xml").anexos[0]
    r = client.get(f"/anexos/{a.id}")
    assert r.headers["Content-Disposition"].startswith("attachment")
    assert r.mimetype == "application/octet-stream"


def test_cabecalhos_seguranca(client):
    r = client.get("/login")
    assert "default-src 'self'" in r.headers["Content-Security-Policy"]
    assert r.headers["X-Frame-Options"] == "DENY"


# ─── avisos por e-mail ──────────────────────────────────────────────────────

def para_quem(emails):
    return [m.para for m in emails]


def test_email_obrigatorio_no_cadastro(client, contas):
    assert "Informe seu e-mail" in cadastrar(client, email="").get_data(as_text=True)
    assert "E-mail inválido" in cadastrar(client, email="nao-e-email").get_data(as_text=True)


def test_conta_sem_email_precisa_cadastrar(client, s, contas):
    ana = usuario(s, "ana")
    ana.email = None
    s.commit()
    entrar(client, "ana", "ana12345")
    r = client.get("/chamados")
    assert r.status_code == 302 and r.location.endswith("/conta/")
    client.post("/conta/", data={"csrf": csrf(client, "/conta/"), "nome": "Ana Souza",
                                 "email": "Ana.Nova@Empresa.com.br"})
    assert usuario(s, "ana").email == "ana.nova@empresa.com.br"
    assert client.get("/chamados").status_code == 200


def test_email_ao_abrir_chamado(client, s, contas, emails):
    entrar(client, "ana", "ana12345")
    abrir(client, s, "Mouse <b>quebrado</b>", setor="T.I", para="carlos", descricao="Não clica\nnada")
    c = chamado(s, "Mouse <b>quebrado</b>")
    assert para_quem(emails) == ["carlos@empresa.com.br"]
    m = emails[0]
    assert m.assunto == f"[Chamado #{c.id}] Novo chamado: Mouse <b>quebrado</b>"
    assert "Ana Souza (Fiscal) abriu um chamado para você." in m.texto
    assert f"/chamados/{c.id}" in m.texto and "Não clica\nnada" in m.texto
    assert "Mouse &lt;b&gt;quebrado&lt;/b&gt;" in m.html and "<b>quebrado</b>" not in m.html


def test_ninguem_recebe_aviso_da_propria_acao(client, s, contas, emails):
    entrar(client, "carlos", "carlos123")          # carlos abre para ele mesmo
    abrir(client, s, "Lembrete meu", setor="T.I", para="carlos")
    assert emails == []


def test_email_de_comentario_vai_para_o_outro_lado(client, s, contas, emails):
    c = chamado(s, "Impressora do fiscal não imprime")   # ana pediu, carlos responsável
    entrar(client, "ana", "ana12345")
    client.post(f"/chamados/{c.id}/comentar", data={"csrf": csrf(client, "/chamados"), "texto": "Alguma novidade?"})
    assert para_quem(emails) == ["carlos@empresa.com.br"]
    assert "Alguma novidade?" in emails[0].texto and "Nova resposta" in emails[0].assunto
    emails.clear()
    sair(client)
    entrar(client, "carlos", "carlos123")
    client.post(f"/chamados/{c.id}/comentar", data={"csrf": csrf(client, "/chamados"), "texto": "Troquei o toner."})
    assert para_quem(emails) == ["ana@empresa.com.br"]


def test_colega_que_comenta_avisa_solicitante_e_responsavel(client, s, contas, emails):
    c = chamado(s, "Erro ao emitir nota fiscal")          # ana pediu, admin responsável
    entrar(client, "carlos", "carlos123")
    client.post(f"/chamados/{c.id}/comentar", data={"csrf": csrf(client, "/chamados"), "texto": "Vou olhar."})
    assert sorted(para_quem(emails)) == ["admin@empresa.com.br", "ana@empresa.com.br"]


def test_email_de_status_e_de_encerramento(client, s, contas, emails):
    c = chamado(s, "Impressora do fiscal não imprime")
    entrar(client, "carlos", "carlos123")
    token = csrf(client, "/chamados")
    client.post(f"/chamados/{c.id}/status", data={"csrf": token, "status": "Aguardando usuário"})
    assert para_quem(emails) == ["ana@empresa.com.br"]
    assert emails[0].assunto.startswith(f"[Chamado #{c.id}] Status: Aguardando usuário")
    emails.clear()
    client.post(f"/chamados/{c.id}/status", data={"csrf": token, "status": "Fechado"})
    assert para_quem(emails) == ["ana@empresa.com.br"]
    assert "Chamado encerrado" in emails[0].assunto
    assert "Avaliar atendimento: " in emails[0].texto and "/avaliacoes" in emails[0].texto


def test_assumir_avisa_solicitante_da_mudanca_de_status(client, s, contas, emails):
    c = chamado(s, "Dúvida sobre saldo de férias")        # bruno pediu, rita responsável, Aberto
    entrar(client, "rita", "rita1234")
    client.post(f"/chamados/{c.id}/assumir", data={"csrf": csrf(client, "/chamados")})
    assert para_quem(emails) == ["bruno@empresa.com.br"]
    assert "Em andamento" in emails[0].assunto


def test_email_ao_encaminhar(client, s, contas, emails):
    c = chamado(s, "Sem acesso à pasta da rede")
    entrar(client, "carlos", "carlos123")
    client.post(f"/chamados/{c.id}/encaminhar", data={"csrf": csrf(client, "/chamados"),
                                                      "setor_destino": "Manutenção",
                                                      "responsavel": usuario(s, "marta").id})
    assert para_quem(emails) == ["marta@empresa.com.br"]
    assert "Carlos Técnico encaminhou este chamado para você (Manutenção)." in emails[0].texto


def test_conta_desativada_nao_recebe_aviso(client, s, contas, emails):
    c = chamado(s, "Impressora do fiscal não imprime")
    ana = usuario(s, "ana")
    ana.ativo = False
    s.commit()
    entrar(client, "carlos", "carlos123")
    client.post(f"/chamados/{c.id}/comentar", data={"csrf": csrf(client, "/chamados"), "texto": "Oi"})
    assert emails == []


def test_link_usa_url_site_configurada(client, s, app, contas, emails):
    app.extensions["email"].url_site = "http://chamados.empresa.local"
    entrar(client, "ana", "ana12345")
    abrir(client, s, "Com link fixo")
    c = chamado(s, "Com link fixo")
    assert f"http://chamados.empresa.local/chamados/{c.id}" in emails[0].texto


def test_config_email_do_ini(tmp_path, monkeypatch):
    from app.config import carregar_config
    ini = tmp_path / "config.ini"
    ini.write_text("[banco]\nurl = sqlite://\n[servidor]\nsecret_key = " + "x" * 40 + "\n"
                   "[email]\nhost = smtp.office365.com\nporta = 587\nseguranca = STARTTLS\n"
                   "usuario = chamados@empresa.com.br\nsenha = a%b\nremetente = chamados@empresa.com.br\n"
                   "url_site = http://chamados.local/\n", encoding="utf-8")
    monkeypatch.setenv("TICKET_MANAGER_CONFIG", str(ini))
    e = carregar_config().email
    assert e.ativo and e.seguranca == "starttls" and e.senha == "a%b" and e.url_site == "http://chamados.local"


def test_mensagem_montada_para_smtp():
    from app.config import ConfigEmail
    from app.correio import Mensagem, _montar
    cfg = ConfigEmail(host="smtp", remetente="chamados@empresa.com.br")
    m = _montar(cfg, Mensagem(para="ana@empresa.com.br", assunto="Título com\nquebra",
                              texto="oi", html="<p>oi</p>"))
    assert m["Subject"] == "Título com quebra"
    assert m["From"] == "Ticket Manager <chamados@empresa.com.br>"
    assert m.get_body(("html",)).get_content().strip() == "<p>oi</p>"


def test_config_email_servidor_interno(tmp_path, monkeypatch):
    import ssl

    from app.config import carregar_config
    from app.correio import _contexto_ssl
    base = "[banco]\nurl = sqlite://\n[servidor]\nsecret_key = " + "x" * 40 + "\n"
    ini = tmp_path / "config.ini"
    monkeypatch.setenv("TICKET_MANAGER_CONFIG", str(ini))

    # relay interno sem senha, porta 25
    ini.write_text(base + "[email]\nhost = mail.empresa.local\nporta = 25\nseguranca = nenhuma\n"
                   "remetente = chamados@empresa.com.br\n", encoding="utf-8")
    e = carregar_config().email
    assert e.ativo and e.porta == 25 and not e.usuario and e.verificar_certificado

    # certificado próprio: arquivo precisa existir
    ini.write_text(base + "[email]\nhost = mail\nremetente = a@b.com\nca_arquivo = /nao/existe.crt\n", encoding="utf-8")
    with pytest.raises(FileNotFoundError):
        carregar_config()

    # verificação desligada
    ini.write_text(base + "[email]\nhost = mail\nremetente = a@b.com\nverificar_certificado = false\n", encoding="utf-8")
    ctx = _contexto_ssl(carregar_config().email)
    assert ctx.verify_mode == ssl.CERT_NONE and not ctx.check_hostname



# ─── ajustes de interface ───────────────────────────────────────────────────

def test_lista_de_usuarios_so_tem_o_botao_editar(client, s, contas):
    entrar(client, "admin", "admin123")
    html = client.get("/usuarios").get_data(as_text=True)
    bruno = usuario(s, "bruno")
    assert f'href="/usuarios/{bruno.id}/editar"' in html
    assert 'name="setor_' not in html and 'name="papel_' not in html                # sem edição na própria lista
    assert "Liberar nova senha</button>" not in html and "Desativar</button>" not in html
    html = client.get(f"/usuarios/{bruno.id}/editar").get_data(as_text=True)
    assert "Liberar nova senha" in html and "Desativar conta" in html and "Apagar usuário" not in html


def test_edicao_nao_muda_o_proprio_perfil(client, s, contas):
    entrar(client, "admin", "admin123")
    salvar_usuarios(client, s, {"admin": ("T.I", True, "usuario")})
    assert usuario(s, "admin").is_admin
    html = client.get(f"/usuarios/{usuario(s, 'admin').id}/editar").get_data(as_text=True)
    assert "Desativar conta" not in html          # a própria conta não tem as ações de senha e conta


def test_apagar_so_usuario_desativado_e_sem_historico(client, s, contas):
    from tests.test_contas import cadastrar_pelo_admin
    entrar(client, "admin", "admin123")
    cadastrar_pelo_admin(client, nome="Cadastro Duplicado", login="duplicado")
    dup = usuario(s, "duplicado")
    token = csrf(client, f"/usuarios/{dup.id}/editar")
    html = client.post(f"/usuarios/{dup.id}/apagar", data={"csrf": token}, follow_redirects=True).get_data(as_text=True)
    assert "Desative a conta antes de apagar" in html and usuario(s, "duplicado") is not None

    client.post(f"/usuarios/{dup.id}/ativo", data={"csrf": token, "ativo": "0"})
    assert "Apagar usuário" in client.get(f"/usuarios/{dup.id}/editar").get_data(as_text=True)
    html = client.post(f"/usuarios/{dup.id}/apagar", data={"csrf": token}, follow_redirects=True).get_data(as_text=True)
    assert "Conta de Cadastro Duplicado apagada." in html and usuario(s, "duplicado") is None

    # quem tem chamados no histórico não é apagado, mesmo desativado
    bruno = usuario(s, "bruno")
    client.post(f"/usuarios/{bruno.id}/ativo", data={"csrf": token, "ativo": "0"})
    html = client.post(f"/usuarios/{bruno.id}/apagar", data={"csrf": token}, follow_redirects=True).get_data(as_text=True)
    assert "tem chamados" in html and usuario(s, "bruno") is not None


def test_tempo_medio_em_dias_horas_e_minutos():
    from app.formatos import duracao
    minuto = 1 / (24 * 60)
    assert duracao(None) == "—"
    assert duracao(0) == "0 min"
    assert duracao(45 * minuto) == "45 min"
    assert duracao(3 / 24 + 20 * minuto) == "3 h 20 min"
    assert duracao(2 / 24) == "2 h"
    assert duracao(2 + 4 / 24 + 10 * minuto) == "2 d 4 h"
    assert duracao(1) == "1 d"


def test_botao_de_tema_e_script_no_cabecalho(client, contas):
    html = client.get("/login").get_data(as_text=True)
    assert "data-alternar-tema" in html and "tema.js" in html
    assert html.index("tema.js") < html.index("style.css")   # aplica o tema antes de desenhar



def test_estaticos_com_versao_para_nao_usar_cache_antigo(client, app):
    import os
    html = client.get("/login").get_data(as_text=True)
    for arquivo in ("style.css", "app.js", "tema.js", "logo.png"):
        mtime = int(os.stat(os.path.join(app.static_folder, arquivo)).st_mtime)
        assert f"/static/{arquivo}?v={mtime}" in html, arquivo
    assert client.get(f"/static/style.css?v=123").status_code == 200   # a versão não atrapalha o arquivo



def test_icone_do_site(client):
    html = client.get("/login").get_data(as_text=True)
    assert "/static/favicon.ico?v=" in html and "/static/icone-180.png?v=" in html
    r = client.get("/favicon.ico")
    assert r.status_code == 200 and r.data[:4] == b"\x00\x00\x01\x00"   # arquivo .ico de verdade



def test_logo_embutido_no_email():
    from app.config import ConfigEmail
    from app.correio import CID_MARCA, Mensagem, _montar
    cfg = ConfigEmail(host="smtp", remetente="chamados@empresa.com.br")
    m = _montar(cfg, Mensagem(para="ana@empresa.com.br", assunto="Teste", texto="oi",
                              html=f'<img src="cid:{CID_MARCA}"><p>oi</p>'))
    tipos = [p.get_content_type() for p in m.walk()]
    assert tipos == ["multipart/alternative", "text/plain", "multipart/related", "text/html", "image/png"]
    imagem = [p for p in m.walk() if p.get_content_type() == "image/png"][0]
    assert imagem["Content-ID"] == f"<{CID_MARCA}>" and imagem.get_payload(decode=True)[:4] == b"\x89PNG"


def test_aviso_usa_logo_embutido(client, s, contas, emails):
    from app.correio import CID_MARCA
    entrar(client, "ana", "ana12345")
    abrir(client, s, "Com logo")
    assert f'src="cid:{CID_MARCA}"' in emails[0].html
    assert "/static/marca.png" in client.get("/chamados").get_data(as_text=True)   # cabeçalho do sistema



def test_cores_css_gerado_de_temas_py():
    from app import temas
    with open(temas.CAMINHO_CSS, encoding="utf-8") as f:
        assert f.read() == temas.gerar_css(), "rode: python -m app.temas"


def test_todas_as_cores_e_intensidades_com_contraste_legivel():
    from app import temas
    for cor in temas.CORES:
        for n in temas.INTENSIDADES:
            for escuro in (False, True):
                assert temas.verificar(cor, n, escuro) == [], (cor, n, escuro)
    for escuro in (False, True):                       # Padrão sobre os fundos do style.css
        v = temas.variaveis_padrao(escuro)
        assert temas.contraste(v["primaria"], v["sobre-primaria"]) >= 4.5
        assert temas.contraste(v["primaria"], "#1b1e22" if escuro else "#ffffff") >= 4.5


def test_intensidade_muda_o_fundo_da_tela():
    from app import temas
    suave, intensa = temas.variaveis("verde", 1, False), temas.variaveis("verde", 5, False)
    assert suave["fundo"] != intensa["fundo"] and suave["borda"] != intensa["borda"]
    css = temas.gerar_css()
    assert ':root[data-cor="verde"][data-intensidade="5"][data-theme="dark"]' in css
    assert css.count("--fundo:") == len(temas.CORES) * len(temas.INTENSIDADES) * 3


def test_seletor_de_cor_na_pagina(client):
    from app import temas
    html = client.get("/login").get_data(as_text=True)
    assert "cores.css?v=" in html and "data-abrir-paleta" in html and "data-intensidade-cor" in html
    for chave, *_ in temas.opcoes():
        assert f'data-cor="{chave}"' in html


# ─── tempo em horário de expediente ─────────────────────────────────────────

def test_segundos_expediente():
    from datetime import datetime as dt
    h = lambda inicio, fim: servicos.segundos_expediente(inicio, fim) / 3600
    # segunda 2026-10-12
    assert h(dt(2026, 10, 12, 9), dt(2026, 10, 12, 11, 30)) == 2.5          # dentro do expediente
    assert h(dt(2026, 10, 12, 17), dt(2026, 10, 13, 8)) == 2                # 17h-18h + 7h-8h (noite não conta)
    assert h(dt(2026, 10, 12, 5), dt(2026, 10, 12, 20)) == 11               # antes e depois não contam
    assert h(dt(2026, 10, 16, 17), dt(2026, 10, 19, 8)) == 2                # sexta 17h → segunda 8h: fim de semana não conta
    assert h(dt(2026, 10, 17, 10), dt(2026, 10, 18, 15)) == 0               # aberto e fechado no fim de semana
    assert h(dt(2026, 10, 12, 7), dt(2026, 10, 16, 18)) == 55               # semana inteira: 5 × 11 h
    assert h(dt(2026, 10, 12, 10), dt(2026, 10, 12, 9)) == 0


def test_painel_tempo_medio_em_expediente(client, s, contas):
    from datetime import datetime as dt
    c = chamado(s, "Trocar teclado")
    c.aberto_em, c.fechado_em = dt(2026, 10, 16, 17), dt(2026, 10, 19, 8)    # sexta 17h → segunda 8h
    for outro in s.scalars(select(Chamado).where(Chamado.status == "Fechado", Chamado.id != c.id)):
        outro.fechado_em = dt(2000, 1, 1)                                      # fora do período
    s.commit()
    d = servicos.painel(s, 365, usuario(s, "admin"))
    assert d["resumo"]["fechados"] == 1 and d["resumo"]["tempo_medio"] == 2
    entrar(client, "admin", "admin123")
    html = client.get("/painel?dias=365").get_data(as_text=True)
    assert "em horário de expediente" in html


def test_painel_do_usuario_comum_mostra_os_chamados_que_ele_abriu(s, contas):
    ana = usuario(s, "ana")
    abertos_ana = s.scalar(select(func.count(Chamado.id)).where(Chamado.solicitante_id == ana.id,
                                                                Chamado.status != "Fechado"))
    assert servicos.painel(s, 30, ana)["resumo"]["abertos"] == abertos_ana


def test_setor_administracao():
    from app.modelos import SETORES
    assert SETORES[0] == "Administração" and SETORES == sorted(SETORES, key=lambda x: x.lower().replace("á", "a"))


# ─── editar comentário (só admin) ───────────────────────────────────────────

def test_admin_edita_comentario_e_fica_registrado(client, s, contas):
    c = chamado(s, "Impressora do fiscal não imprime")
    m = next(x for x in c.comentarios if not x.evento)
    url = f"/chamados/{c.id}/comentarios/{m.id}/editar"

    entrar(client, "carlos", "carlos123")              # atende o chamado, mas não é admin
    assert "editar-comentario" not in client.get(f"/chamados/{c.id}").get_data(as_text=True)
    assert client.post(url, data={"csrf": csrf(client, f"/chamados/{c.id}"), "texto": "x"}).status_code == 403
    sair(client)

    entrar(client, "admin", "admin123")
    html = client.get(f"/chamados/{c.id}").get_data(as_text=True)
    assert 'class="editar-comentario"' in html
    r = client.post(url, data={"csrf": csrf(client, f"/chamados/{c.id}"), "texto": "  Texto corrigido.  "})
    assert r.location.endswith(f"#m{m.id}")
    s.expire_all()
    m = s.get(Comentario, m.id)
    assert m.texto == "Texto corrigido." and m.editado_por.login == "admin" and m.editado_em
    html = client.get(f"/chamados/{c.id}").get_data(as_text=True)
    assert "editado por Administrador do Sistema" in html

    html = client.post(url, data={"csrf": csrf(client, f"/chamados/{c.id}"), "texto": "   "},
                       follow_redirects=True).get_data(as_text=True)
    assert "não pode ficar vazio" in html and s.get(Comentario, m.id).texto == "Texto corrigido."


def test_registro_automatico_nao_e_editavel(client, s, contas):
    entrar(client, "admin", "admin123")
    c = chamado(s, "Impressora do fiscal não imprime")
    servicos.alterar_status(s, c, "Aguardando usuário", usuario(s, "admin"))   # gera um registro automático
    evento = next(x for x in chamado(s, "Impressora do fiscal não imprime").comentarios if x.evento)
    r = client.post(f"/chamados/{c.id}/comentarios/{evento.id}/editar",
                    data={"csrf": csrf(client, f"/chamados/{c.id}"), "texto": "outro"}, follow_redirects=True)
    assert "não podem ser editados" in r.get_data(as_text=True)
    outro = chamado(s, "Trocar teclado")                 # comentário de outro chamado: 404
    assert client.post(f"/chamados/{outro.id}/comentarios/{evento.id}/editar",
                       data={"csrf": csrf(client, f"/chamados/{c.id}"), "texto": "x"}).status_code == 404


# ─── mudar status direto na lista ───────────────────────────────────────────

def test_mudar_status_direto_na_lista(client, s, contas, emails):
    c = chamado(s, "Impressora do fiscal não imprime")            # T.I, aberto, responsável carlos
    entrar(client, "carlos", "carlos123")
    html = client.get("/chamados?status=abertos&pagina=1").get_data(as_text=True)
    assert f'aria-label="Status do chamado #{c.id}"' in html
    rh = chamado(s, "Dúvida sobre saldo de férias")                # RH: carlos não atende
    assert f'aria-label="Status do chamado #{rh.id}"' not in html

    voltar = "/chamados?status=abertos&pagina=1"
    r = client.post(f"/chamados/{c.id}/status",
                    data={"csrf": csrf(client, "/chamados"), "status": "Em andamento", "voltar": voltar})
    assert r.location.endswith(f"{voltar}#c{c.id}")                # volta para a mesma lista, na linha
    assert chamado(s, "Impressora do fiscal não imprime").status == "Em andamento"
    assert any("Em andamento" in m.assunto for m in emails)         # mesmo aviso de quando muda pelo chamado
    assert f"Chamado #{c.id}: status alterado para Em andamento." in client.get(r.location).get_data(as_text=True)

    # destino de fora do sistema é ignorado
    r = client.post(f"/chamados/{c.id}/status",
                    data={"csrf": csrf(client, "/chamados"), "status": "Aberto", "voltar": "//outro.site/x"})
    assert r.location.endswith(f"/chamados/#c{c.id}") or r.location.endswith(f"/chamados#c{c.id}")


def test_quem_so_abriu_o_chamado_nao_muda_status_na_lista(client, s, contas):
    entrar(client, "ana", "ana12345")
    html = client.get("/chamados").get_data(as_text=True)
    assert 'class="status-rapido' not in html and "badge-aberto" in html


# ─── painel com gráficos ────────────────────────────────────────────────────

def test_indicadores_do_painel(client, s, contas):
    from datetime import datetime as dt, timedelta as td
    admin = usuario(s, "admin")
    agora_ = dt.now().replace(microsecond=0)
    # dois encerrados neste período e um no anterior, para a comparação
    cs = s.scalars(select(Chamado).order_by(Chamado.id)).all()
    for c, (abre, fecha) in zip(cs[:3], [(3, 2), (5, 4), (45, 40)]):
        c.status, c.aberto_em, c.fechado_em = "Fechado", agora_ - td(days=abre), agora_ - td(days=fecha)
    for c in cs[3:]:
        if c.status == "Fechado":
            c.fechado_em = dt(2000, 1, 1)
    s.commit()
    d = servicos.painel(s, 30, admin)
    r = d["resumo"]
    assert r["fechados"] == 2 and r["fechados_var"] == 100.0          # 2 agora x 1 no período anterior
    assert d["agrupamento"] == "semana" and len(d["serie"]) in (5, 6)
    assert sum(b["encerrados"] for b in d["serie"]) == 2
    assert [rot for rot, _n in d["idade"]] == ["Até 4 h", "4 h a 1 dia", "1 a 3 dias", "Mais de 3 dias"]
    assert sum(n for _r, n in d["idade"]) == r["abertos"] == sum(n for _s, n in d["por_status"])
    assert sum(sum(v) for _n, v in d["setor_prioridade"]) == r["abertos"] and d["por_destino"]
    assert [dia for dia, _v in d["calor"]["dias"]][:5] == ["Seg", "Ter", "Qua", "Qui", "Sex"]
    assert servicos.painel(s, 365, admin)["agrupamento"] == "mês"

    entrar(client, "admin", "admin123")
    html = client.get("/painel?dias=90").get_data(as_text=True)
    for trecho in ("Abertos x encerrados por semana", "Em aberto por status", "Há quanto tempo estão em aberto",
                   "Quando os chamados chegam", "Avaliação do atendimento", "Por responsável", 'class="spark"',
                   '<svg class="grafico-svg"', '<svg class="rosca"'):
        assert trecho in html, trecho


def test_pdf_do_relatorio_com_graficos(client, s, contas):
    from app import exportacao, graficos_pdf
    from datetime import date as d_, datetime as dt
    from reportlab.lib.styles import getSampleStyleSheet
    d = servicos.indicadores(s, dt(2000, 1, 1), dt(2100, 1, 1), usuario(s, "admin"))
    partes = graficos_pdf.resumo(d, getSampleStyleSheet(), 270)
    assert len(partes) >= 6
    com = exportacao.gerar_pdf([], d_(2026, 1, 1), d_(2026, 1, 31), d, "Todos os setores").getvalue()
    sem = exportacao.gerar_pdf([], d_(2026, 1, 1), d_(2026, 1, 31)).getvalue()
    assert com[:4] == b"%PDF" and len(com) > len(sem) + 3000



# ─── lembrete de avaliação ──────────────────────────────────────────────────

def test_lembrete_de_avaliacao(client, s, app, contas):
    from datetime import datetime as dt, timedelta as td
    from app import notificacoes
    agora_ = dt.now().replace(microsecond=0)
    for c in s.scalars(select(Chamado).where(Chamado.status == "Fechado")):
        c.avaliacao = "Bom"                                   # zera o cenário
    s.commit()
    a, b = chamado(s, "Impressora do fiscal não imprime"), chamado(s, "Erro ao emitir nota fiscal")
    a.status, a.fechado_em, a.avaliacao = "Fechado", agora_ - td(days=2), None     # ana, já pode lembrar
    b.status, b.fechado_em, b.avaliacao = "Fechado", agora_ - td(hours=5), None    # ana, encerrado agora há pouco
    s.commit()

    pend = servicos.lembretes_de_avaliacao(s)
    ana = usuario(s, "ana")
    assert list(pend) == [ana] and [c.id for c in pend[ana]] == [a.id]
    with app.test_request_context():
        msg = notificacoes.mensagem_lembrete_avaliacao(ana, pend[ana])
    assert msg.para == "ana@empresa.com.br" and f"#{a.id}" in msg.assunto and "/avaliacoes" in msg.texto

    servicos.registrar_lembretes(s, pend[ana])
    assert servicos.lembretes_de_avaliacao(s) == {}           # espera 3 dias para o próximo
    a = chamado(s, "Impressora do fiscal não imprime")
    a.ultimo_lembrete_em = agora_ - td(days=4)
    s.commit()
    assert [c.id for c in servicos.lembretes_de_avaliacao(s)[usuario(s, "ana")]] == [a.id]
    servicos.registrar_lembretes(s, [a])
    a = chamado(s, "Impressora do fiscal não imprime")
    a.ultimo_lembrete_em = agora_ - td(days=10)
    s.commit()
    assert servicos.lembretes_de_avaliacao(s) == {}           # no máximo 2 lembretes

    # aviso fixo no topo enquanto houver pendente
    entrar(client, "ana", "ana12345")
    assert "★ Avaliar" in client.get("/conta/").get_data(as_text=True)
    token = csrf(client, "/avaliacoes")
    for c in (a, chamado(s, "Erro ao emitir nota fiscal")):
        client.post(f"/chamados/{c.id}/avaliar", data={"csrf": token, "avaliacao": "Bom"})
    assert "★ Avaliar" not in client.get("/conta/").get_data(as_text=True)
