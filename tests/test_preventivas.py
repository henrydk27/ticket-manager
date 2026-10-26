"""Testes da revisão preventiva dos ativos."""

import io
import os
from datetime import date, datetime, timedelta

from openpyxl import load_workbook
from sqlalchemy import select

from app import preventivas
from app.modelos import PLANOS_PADRAO, Chamado, Preventiva
from tests.conftest import csrf, entrar
from tests.test_app import sair, usuario
from tests.test_inventario import cadastrar, editar, eq


def em_uso(client, s, patrimonio="000123", **campos):
    """Desktop em uso no Fiscal (cadastrado pelo admin já logado)."""
    cadastrar(client, patrimonio=patrimonio, numero_serie=f"S{patrimonio}", hostname="", ip="", mac="",
              situacao="Em uso", setor="Fiscal", **campos)
    return eq(s, patrimonio)


def registrar(client, e, resultado="OK", data=None, **extra):
    plano = PLANOS_PADRAO[e.tipo][1]
    dados = {"csrf": csrf(client, f"/preventivas/equipamento/{e.id}"), "data": (data or date.today()).isoformat()}
    for i, _item in enumerate(plano):
        dados[f"r{i}"] = resultado
    dados.update(extra)
    return client.post(f"/preventivas/equipamento/{e.id}", data=dados, content_type="multipart/form-data",
                       follow_redirects=True)


def test_quem_acessa_as_preventivas(client, s, contas):
    for login, senha, lista, planos in (("ana", "ana12345", 403, 403), ("marta", "marta123", 403, 403),
                                        ("carlos", "carlos123", 200, 403), ("admin", "admin123", 200, 200)):
        entrar(client, login, senha)
        assert client.get("/preventivas/").status_code == lista, login
        assert client.get("/preventivas/planos").status_code == planos, login
        menu = client.get("/chamados").get_data(as_text=True)
        assert (">Preventivas<" in menu) == (lista == 200), login
        sair(client)


def test_proxima_preventiva_pelo_plano_do_tipo(client, s, contas):
    entrar(client, "admin", "admin123")
    e = em_uso(client, s)
    situ = preventivas.situacao(e, preventivas.planos(s))
    assert situ.meses == 6 and situ.ultima is None
    assert situ.proxima == preventivas.somar_meses(e.criado_em.date(), 6) and situ.estado == "em_dia"

    e.criado_em = datetime.now() - timedelta(days=200)               # cadastrado há mais de 6 meses
    s.commit()
    assert preventivas.situacao(e, preventivas.planos(s)).estado == "atrasada"
    e.criado_em = datetime.now() - timedelta(days=175)
    s.commit()
    assert preventivas.situacao(e, preventivas.planos(s)).estado == "vence"

    editar(client, eq(s, "000123"), situacao="Em uso", preventiva_meses="12")   # periodicidade própria
    assert preventivas.situacao(eq(s, "000123"), preventivas.planos(s)).meses == 12
    editar(client, eq(s, "000123"), situacao="Em uso", preventiva_meses="0")    # não faz preventiva
    assert preventivas.situacao(eq(s, "000123"), preventivas.planos(s)).estado == "sem_plano"
    assert "escolha uma das opções" in editar(client, eq(s, "000123"), situacao="Em uso",
                                              preventiva_meses="7").get_data(as_text=True)
    editar(client, eq(s, "000123"), situacao="Em estoque", preventiva_meses="")  # fora de uso: não conta
    assert preventivas.situacao(eq(s, "000123"), preventivas.planos(s)).estado == "sem_plano"
    assert preventivas.somar_meses(date(2026, 8, 31), 6) == date(2027, 2, 28)


def test_registrar_preventiva_ok_recalcula_a_proxima(client, s, app, contas):
    entrar(client, "admin", "admin123")
    e = em_uso(client, s)
    e.criado_em = datetime.now() - timedelta(days=400)
    s.commit()
    sair(client)

    entrar(client, "carlos", "carlos123")                          # atende a T.I
    lista = client.get("/preventivas/").get_data(as_text=True)
    assert "000123" in lista and "Atrasada" in lista
    html = client.get(f"/preventivas/equipamento/{e.id}").get_data(as_text=True)
    assert "Limpeza interna e das saídas de ar" in html and "Windows Update" in html

    assert "não pode ser no futuro" in registrar(client, e, data=date.today() + timedelta(days=1)).get_data(as_text=True)
    r = client.post(f"/preventivas/equipamento/{e.id}", data={"csrf": csrf(client, "/preventivas/"),
                    "data": date.today().isoformat()}, follow_redirects=True)
    assert "Marque o resultado de" in r.get_data(as_text=True)

    html = registrar(client, e, pecas="pasta térmica", observacoes="Tudo certo.",
                     anexo=(io.BytesIO(b"\x89PNG foto"), "foto.png")).get_data(as_text=True)
    assert "Preventiva do patrimônio 000123 registrada." in html
    e = eq(s, "000123")
    p = e.preventivas[0]
    assert p.tecnico.login == "carlos" and p.pecas == "pasta térmica" and p.chamado_id is None
    assert p.contar("OK") == len(p.itens) == 8 and p.anexo_nome == "foto.png"
    assert e.historico[0].descricao.startswith("Preventiva de ")
    situ = preventivas.situacao(e, preventivas.planos(s))
    assert situ.estado == "em_dia" and situ.proxima == preventivas.somar_meses(date.today(), 6)
    assert client.get(f"/preventivas/{p.id}/anexo").status_code == 200
    arquivo = os.path.join(app.config["ANEXOS_PASTA"], p.anexo_arquivo)
    assert os.path.exists(arquivo)
    sair(client)

    # apagar o equipamento leva junto as preventivas e o arquivo
    entrar(client, "admin", "admin123")
    client.post(f"/inventario/{e.id}/apagar", data={"csrf": csrf(client, "/inventario/")})
    assert not os.path.exists(arquivo) and s.scalar(select(Preventiva.id)) is None


def test_item_com_problema_abre_chamado_para_a_ti(client, s, contas, emails):
    entrar(client, "admin", "admin123")
    e = em_uso(client, s, usuario_id=usuario(s, "ana").id)
    sair(client)
    entrar(client, "carlos", "carlos123")
    plano = PLANOS_PADRAO["Desktop"][1]
    dados = {"csrf": csrf(client, f"/preventivas/equipamento/{e.id}"), "data": date.today().isoformat()}
    for i, _item in enumerate(plano):
        dados[f"r{i}"] = "OK"
    dados["r4"], dados["n4"] = "Problema", "disco com setores defeituosos"
    dados["r2"] = "Corrigido"
    html = client.post(f"/preventivas/equipamento/{e.id}", data=dados, follow_redirects=True).get_data(as_text=True)
    assert "chamado #" in html and "aberto para a T.I" in html
    p = eq(s, "000123").preventivas[0]
    assert p.contar("Problema") == 1 and p.contar("Corrigido") == 1
    c = s.get(Chamado, p.chamado_id)
    assert c.setor_destino == "T.I" and c.responsavel.login == "carlos" and c.solicitante.login == "carlos"
    assert "000123" in c.titulo and "disco com setores defeituosos" in c.descricao and "Ana Souza" in c.descricao


def test_planos_editaveis_pelo_admin(client, s, contas):
    entrar(client, "admin", "admin123")
    html = client.get("/preventivas/planos").get_data(as_text=True)
    assert "Roletes de tração" in html
    tipos = list(preventivas.planos(s))
    dados = {"csrf": csrf(client, "/preventivas/planos")}
    for i, t in enumerate(tipos):
        p = preventivas.planos(s)[t]
        dados[f"meses_{i}"], dados[f"itens_{i}"] = str(p.meses), "\n".join(p.itens)
    i = tipos.index("Impressora")
    dados[f"meses_{i}"], dados[f"itens_{i}"] = "2", "Limpeza\n\n  Troca do toner  \n"
    assert "Planos de preventiva salvos." in client.post("/preventivas/planos", data=dados,
                                                         follow_redirects=True).get_data(as_text=True)
    assert preventivas.planos(s)["Impressora"].meses == 2
    assert preventivas.planos(s)["Impressora"].itens == ["Limpeza", "Troca do toner"]

    dados[f"meses_{i}"], dados[f"itens_{i}"] = "3", ""                       # meses sem itens: erro
    html = client.post("/preventivas/planos", data=dados, follow_redirects=True).get_data(as_text=True)
    assert "Nada foi salvo" in html and preventivas.planos(s)["Impressora"].meses == 2


def test_relatorio_excel_e_resumo_semanal(client, s, app, contas):
    from app import notificacoes
    entrar(client, "admin", "admin123")
    e = em_uso(client, s)
    registrar(client, e)
    atrasado = em_uso(client, s, patrimonio="000124")
    atrasado.criado_em = datetime.now() - timedelta(days=400)
    s.commit()

    hoje = date.today()
    r = client.get(f"/preventivas/relatorio?de={(hoje - timedelta(days=30)).isoformat()}&ate={hoje.isoformat()}")
    assert r.status_code == 200
    linhas = list(load_workbook(io.BytesIO(r.data)).active.iter_rows(values_only=True))
    assert linhas[0][:3] == ("Data", "Patrimônio", "Tipo") and [l[1] for l in linhas[1:]] == ["000123"]
    assert client.get("/preventivas/relatorio?de=2026-10-10&ate=2026-10-01").status_code == 302

    app.extensions["email"].avisos_cpd = "cpd@empresa.com.br"
    with app.test_request_context():
        msg = notificacoes.mensagem_semanal(preventivas.pendentes(s))
    assert msg.para == "cpd@empresa.com.br" and "1 preventiva pendente" in msg.assunto
    assert "000124" in msg.texto and "000123" not in msg.texto
    with app.test_request_context():
        assert notificacoes.mensagem_semanal([]) is None


def test_registrar_pela_ficha_volta_para_a_ficha(client, s, contas):
    entrar(client, "admin", "admin123")
    e = em_uso(client, s)
    ficha = f"/inventario/{e.id}"
    html = client.get(ficha).get_data(as_text=True)
    assert f"/preventivas/equipamento/{e.id}?voltar=/inventario/{e.id}" in html
    html = client.get(f"/preventivas/equipamento/{e.id}?voltar={ficha}").get_data(as_text=True)
    assert "Marcar todos como OK" in html and "Ficha do equipamento" in html
    dados = {"csrf": csrf(client, ficha), "data": date.today().isoformat(), "voltar": ficha}
    for i, _item in enumerate(PLANOS_PADRAO["Desktop"][1]):
        dados[f"r{i}"] = "OK"
    r = client.post(f"/preventivas/equipamento/{e.id}", data=dados)
    assert r.location.endswith(ficha)
    dados["voltar"] = "https://outro.site/"
    assert client.post(f"/preventivas/equipamento/{e.id}", data=dados).location.endswith("/preventivas/")
